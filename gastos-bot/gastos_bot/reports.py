"""Textos que envía el bot (HTML de Telegram) y exportación a Excel."""

from __future__ import annotations

import calendar
import io
from collections import defaultdict
from datetime import date
from html import escape
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import FiscalData
from .db import Expense

TELEGRAM_LIMIT = 4000

MONTH_NAMES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)  # fmt: skip

CATEGORY_NAMES = {
    "alimentos": "Alimentos",
    "supermercado": "Supermercado",
    "gasolina": "Gasolina",
    "transporte": "Transporte",
    "hospedaje": "Hospedaje",
    "servicios": "Servicios",
    "salud": "Salud",
    "oficina": "Oficina",
    "tecnologia": "Tecnología",
    "hogar": "Hogar",
    "entretenimiento": "Entretenimiento",
    "otros": "Otros",
}
CATEGORY_EMOJI = {
    "alimentos": "🍽️",
    "supermercado": "🛒",
    "gasolina": "⛽",
    "transporte": "🚕",
    "hospedaje": "🏨",
    "servicios": "💡",
    "salud": "💊",
    "oficina": "📎",
    "tecnologia": "💻",
    "hogar": "🏠",
    "entretenimiento": "🎬",
    "otros": "📦",
}
STATUS_NAMES = {"pendiente": "Pendiente", "facturado": "Facturado", "no_requiere": "No requiere"}
STATUS_ICONS = {"pendiente": "⏳", "facturado": "✅", "no_requiere": "➖"}


def esc(value: object) -> str:
    return escape(str(value), quote=False)


def plural(n: int, word: str) -> str:
    return f"{n} {word}" if n == 1 else f"{n} {word}s"


def money(amount: float | None, currency: str = "MXN") -> str:
    if amount is None:
        return "sin total"
    text = f"${amount:,.2f}"
    return text if currency == "MXN" else f"{text} {currency}"


def month_label(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month - 1]} {year}"


def category_label(category: str) -> str:
    return f"{CATEGORY_EMOJI.get(category, '📦')} {CATEGORY_NAMES.get(category, category)}"


def _invoice_data_lines(e: Expense) -> list[str]:
    """Datos que piden los portales para facturar este ticket."""
    lines = []
    ids = []
    if e.rfc:
        ids.append(f"RFC <code>{esc(e.rfc)}</code>")
    if e.folio:
        ids.append(f"Folio <code>{esc(e.folio)}</code>")
    if ids:
        lines.append(" · ".join(ids))
    if e.invoice_portal:
        lines.append(f"Portal: {esc(e.invoice_portal)}")
    for item in e.invoice_fields:
        lines.append(f"{esc(item.get('label', ''))}: <code>{esc(item.get('value', ''))}</code>")
    return lines


def expense_detail(e: Expense, title: str | None = None) -> str:
    lines = [title or f"<b>Gasto #{e.id}</b>", f"🏪 {esc(e.merchant or 'Comercio sin nombre')}"]
    if e.legal_name and e.legal_name != e.merchant:
        lines.append(f"     {esc(e.legal_name)}")
    lines.append(f"📅 {e.expense_date:%d/%m/%Y} · <b>{money(e.total, e.currency)}</b>")
    category = category_label(e.category)
    lines.append(f"{category} · {esc(e.description)}" if e.description else category)
    if e.payment_method:
        lines.append(f"💳 {esc(e.payment_method)}")
    if e.user_name:
        lines.append(f"👤 {esc(e.user_name)}")
    lines.append(f"🧾 Factura: {STATUS_ICONS[e.invoice_status]} {STATUS_NAMES[e.invoice_status].lower()}")
    if e.invoice_status == "facturado" and e.cfdi_uuid:
        lines.append(f"     UUID <code>{esc(e.cfdi_uuid)}</code>")
    elif e.invoice_status == "pendiente":
        lines.extend(f"     {line}" for line in _invoice_data_lines(e))
    if e.notes:
        lines.append(f"📝 {esc(e.notes)}")
    return "\n".join(lines)


def expense_line(e: Expense) -> str:
    name = esc(e.merchant or e.description or "Comercio sin nombre")
    return (
        f"#{e.id} · {e.expense_date:%d/%m} · {name} · "
        f"{money(e.total, e.currency)} {STATUS_ICONS[e.invoice_status]}"
    )


def recent_text(expenses: list[Expense]) -> str:
    if not expenses:
        return "Aún no hay gastos registrados. Envía la foto de un ticket para empezar."
    lines = ["🧾 <b>Últimos gastos</b>", *(expense_line(e) for e in expenses)]
    lines += ["", f"Detalle y foto de un gasto: /ver {expenses[0].id}"]
    return "\n".join(lines)


def summary_text(expenses: list[Expense], period: str) -> str:
    if not expenses:
        return f"No hay gastos registrados en {period}."

    mxn = [e for e in expenses if e.total is not None and e.currency == "MXN"]
    other = [e for e in expenses if e.total is not None and e.currency != "MXN"]
    missing = [e for e in expenses if e.total is None]
    total = sum(e.total for e in mxn)

    lines = [f"📊 <b>Resumen de {period}</b>", f"Total: <b>{money(total)}</b> en {plural(len(mxn), 'gasto')}"]

    by_category: dict[str, float] = defaultdict(float)
    by_person: dict[str, list[float]] = defaultdict(list)
    for e in mxn:
        by_category[e.category] += e.total
        by_person[e.user_name or "Sin nombre"].append(e.total)

    if by_category:
        lines += ["", "<b>Por categoría</b>"]
        for category, amount in sorted(by_category.items(), key=lambda kv: -kv[1]):
            share = f" ({amount / total:.0%})" if total else ""
            lines.append(f"{category_label(category)}: {money(amount)}{share}")

    if len(by_person) > 1:
        lines += ["", "<b>Por persona</b>"]
        for person, amounts in sorted(by_person.items(), key=lambda kv: -sum(kv[1])):
            lines.append(f"👤 {esc(person)}: {money(sum(amounts))} ({plural(len(amounts), 'gasto')})")

    lines += ["", "<b>Facturación</b>"]
    for status, label in (("facturado", "Facturados"), ("pendiente", "Pendientes"), ("no_requiere", "No requieren")):
        group = [e for e in mxn if e.invoice_status == status]
        if group:
            lines.append(f"{STATUS_ICONS[status]} {label}: {len(group)} · {money(sum(e.total for e in group))}")

    if other:
        by_currency: dict[str, float] = defaultdict(float)
        for e in other:
            by_currency[e.currency] += e.total
        amounts = ", ".join(money(v, c) for c, v in sorted(by_currency.items()))
        lines += ["", f"💱 Además, en otras monedas: {amounts}"]
    if missing:
        ids = ", ".join(f"#{e.id}" for e in missing)
        lines += [
            "",
            f"⚠️ Sin total: {ids}. Corrígelos con /editar {missing[0].id} total 123.45",
        ]
    return "\n".join(lines)


def _pending_entry(e: Expense) -> str:
    who = f" · {esc(e.user_name)}" if e.user_name else ""
    lines = [
        f"#{e.id} · {e.expense_date:%d/%m} · <b>{esc(e.merchant or 'Comercio sin nombre')}</b> · "
        f"{money(e.total, e.currency)}{who}"
    ]
    lines.extend(f"     {line}" for line in _invoice_data_lines(e))
    return "\n".join(lines)


def _days_left_text(today: date) -> str:
    last_day = calendar.monthrange(today.year, today.month)[1]
    days_left = last_day - today.day
    if days_left == 0:
        return "hoy es el último día del mes"
    return f"quedan {plural(days_left, 'día')} del mes"


def pending_text(expenses: list[Expense], today: date) -> str:
    if not expenses:
        return "✅ No hay tickets pendientes de facturar."
    month_start = today.replace(day=1)
    current = [e for e in expenses if e.expense_date >= month_start]
    older = [e for e in expenses if e.expense_date < month_start]
    total = sum(e.total for e in expenses if e.total is not None and e.currency == "MXN")

    lines = [f"⏳ <b>{plural(len(expenses), 'ticket')} por facturar</b> · {money(total)}"]
    if current:
        lines += ["", f"<b>{month_label(today.year, today.month).capitalize()}</b> — {_days_left_text(today)}."]
        lines += [_pending_entry(e) for e in current]
    if older:
        lines += [
            "",
            "<b>Meses anteriores</b> — muchos comercios solo facturan dentro del mes de la "
            "compra; revisa si todavía se puede.",
        ]
        lines += [_pending_entry(e) for e in older]
    lines += [
        "",
        f"Ya lo facturaste: /facturado {expenses[0].id} (o envía aquí el XML de la factura)",
        f"No se va a facturar: /nofacturar {expenses[0].id}",
        "Tus datos para los portales: /datosfiscales",
    ]
    return "\n".join(lines)


def reminder_text(expenses: list[Expense], today: date) -> str:
    lines = [
        "⏰ <b>Recordatorio de facturación</b>",
        f"Tienen {plural(len(expenses), 'ticket')} de {month_label(today.year, today.month)} sin "
        f"facturar y {_days_left_text(today)}:",
        "",
        *(_pending_entry(e) for e in expenses),
        "",
        f"Cuando los facturen: /facturado {expenses[0].id} (o envíen aquí el XML)",
    ]
    return "\n".join(lines)


def fiscal_text(fiscal: FiscalData) -> str:
    if fiscal.is_empty():
        return (
            "No hay datos fiscales configurados. Agrega FISCAL_RFC, FISCAL_NOMBRE, FISCAL_CP, "
            "FISCAL_REGIMEN, FISCAL_USO_CFDI y FISCAL_EMAIL a la configuración del bot."
        )
    rows = (
        ("RFC", fiscal.rfc),
        ("Nombre o razón social", fiscal.nombre),
        ("Código postal", fiscal.codigo_postal),
        ("Régimen fiscal", fiscal.regimen),
        ("Uso de CFDI", fiscal.uso_cfdi),
        ("Correo", fiscal.email),
    )
    lines = ["🧾 <b>Datos para facturar</b> (toca un dato para copiarlo)"]
    lines += [f"{label}: <code>{esc(value)}</code>" for label, value in rows if value]
    return "\n".join(lines)


def split_message(text: str, limit: int = TELEGRAM_LIMIT) -> list[str]:
    """Parte un texto largo por renglones para no pasar el límite de Telegram."""
    chunks, current = [], ""
    for line in text.split("\n"):
        candidate = f"{current}\n{line}" if current else line
        if len(candidate) > limit and current:
            chunks.append(current)
            current = line
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


_XLSX_COLUMNS = (
    ("ID", 6),
    ("Fecha", 12),
    ("Comercio", 24),
    ("Razón social", 32),
    ("RFC", 16),
    ("Categoría", 16),
    ("Descripción", 36),
    ("Subtotal", 12),
    ("IVA", 12),
    ("Total", 12),
    ("Moneda", 8),
    ("Forma de pago", 18),
    ("Folio", 16),
    ("Factura", 13),
    ("UUID", 38),
    ("Portal de facturación", 28),
    ("Datos para facturar", 36),
    ("Registrado por", 16),
    ("Registrado el", 17),
    ("Notas", 36),
)


def build_xlsx(expenses: list[Expense], tz: ZoneInfo) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"
    ws.append([name for name, _ in _XLSX_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for index, (_, width) in enumerate(_XLSX_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for e in expenses:
        ws.append(
            [
                e.id,
                e.expense_date,
                e.merchant,
                e.legal_name,
                e.rfc,
                CATEGORY_NAMES.get(e.category, e.category),
                e.description,
                e.subtotal,
                e.iva,
                e.total,
                e.currency,
                e.payment_method,
                e.folio,
                STATUS_NAMES[e.invoice_status],
                e.cfdi_uuid,
                e.invoice_portal,
                "; ".join(f"{f.get('label', '')}: {f.get('value', '')}" for f in e.invoice_fields),
                e.user_name,
                e.created_at.astimezone(tz).replace(tzinfo=None),
                e.notes,
            ]
        )
        row = ws.max_row
        ws.cell(row, 2).number_format = "dd/mm/yyyy"
        for column in (8, 9, 10):
            ws.cell(row, column).number_format = "#,##0.00"
        ws.cell(row, 19).number_format = "dd/mm/yyyy hh:mm"
        for cell in ws[row]:
            # Un texto leído del ticket que empiece con "=" no debe convertirse en fórmula.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
