import io
from datetime import date

from openpyxl import load_workbook

from gastos_bot.config import FiscalData
from gastos_bot.reports import (
    build_xlsx,
    expense_detail,
    fiscal_text,
    money,
    pending_text,
    recent_text,
    reminder_text,
    split_message,
    summary_text,
)
from tests.conftest import TZ, add_expense

TODAY = date(2026, 10, 14)


def test_money():
    assert money(1234.5) == "$1,234.50"
    assert money(10, "USD") == "$10.00 USD"
    assert money(None) == "sin total"


def test_expense_detail_escapes_html_and_shows_invoice_data(db):
    e = add_expense(
        db,
        merchant="Tacos <El Güero> & Cía",
        rfc="CCO8605231N4",
        folio="A-1",
        invoice_portal="facturacion.oxxo.com",
        invoice_fields=[{"label": "Web ID", "value": "<99>"}],
    )
    text = expense_detail(e)
    assert "Tacos &lt;El Güero&gt; &amp; Cía" in text
    assert "RFC <code>CCO8605231N4</code>" in text
    assert "Folio <code>A-1</code>" in text
    assert "Portal: facturacion.oxxo.com" in text
    assert "Web ID: <code>&lt;99&gt;</code>" in text
    assert "⏳ pendiente" in text


def test_expense_detail_invoiced_shows_uuid_instead_of_portal(db):
    e = add_expense(db, invoice_status="facturado", cfdi_uuid="AAA", invoice_portal="x.com")
    text = expense_detail(e, title="Título")
    assert text.startswith("Título")
    assert "UUID <code>AAA</code>" in text
    assert "x.com" not in text


def test_summary_text(db):
    expenses = [
        add_expense(db, total=600.0, category="gasolina", user_name="Ana", invoice_status="facturado"),
        add_expense(db, total=300.0, category="supermercado", user_name="Beto"),
        add_expense(db, total=100.0, category="gasolina", user_name="Ana", invoice_status="no_requiere"),
        add_expense(db, total=50.0, currency="USD", user_name="Ana"),
        add_expense(db, total=None, user_name="Beto"),
    ]
    text = summary_text(expenses, "octubre 2026")
    assert "Total: <b>$1,000.00</b> en 3 gastos" in text
    assert "⛽ Gasolina: $700.00 (70%)" in text
    assert "🛒 Supermercado: $300.00 (30%)" in text
    assert "👤 Ana: $700.00 (2 gastos)" in text
    assert "👤 Beto: $300.00 (1 gasto)" in text
    assert "✅ Facturados: 1 · $600.00" in text
    assert "⏳ Pendientes: 1 · $300.00" in text
    assert "➖ No requieren: 1 · $100.00" in text
    assert "$50.00 USD" in text
    assert f"Sin total: #{expenses[4].id}" in text


def test_summary_without_expenses():
    assert summary_text([], "octubre 2026") == "No hay gastos registrados en octubre 2026."


def test_pending_text_groups_by_month(db):
    old = add_expense(db, merchant="Pemex", expense_date=date(2026, 9, 28), total=800.0)
    new = add_expense(db, merchant="OXXO", expense_date=date(2026, 10, 3), total=200.0, rfc="CCO8605231N4")
    text = pending_text([old, new], TODAY)
    assert "2 tickets por facturar</b> · $1,000.00" in text
    assert "<b>Octubre 2026</b> — quedan 17 días del mes." in text
    assert text.index("OXXO") < text.index("Meses anteriores") < text.index("Pemex")
    assert f"/facturado {old.id}" in text


def test_pending_text_last_day_and_empty(db):
    e = add_expense(db, expense_date=date(2026, 10, 31))
    assert "hoy es el último día del mes" in pending_text([e], date(2026, 10, 31))
    assert pending_text([], TODAY).startswith("✅")


def test_reminder_text(db):
    e = add_expense(db)
    text = reminder_text([e], date(2026, 10, 25))
    assert "1 ticket de octubre 2026" in text
    assert "quedan 6 días del mes" in text
    assert f"#{e.id}" in text


def test_recent_text(db):
    assert "Aún no hay gastos" in recent_text([])
    e = add_expense(db)
    assert f"#{e.id} · 03/10 · OXXO · $245.50 ⏳" in recent_text([e])


def test_fiscal_text():
    assert "No hay datos fiscales" in fiscal_text(FiscalData())
    text = fiscal_text(FiscalData(rfc="XAXX010101000", codigo_postal="64000"))
    assert "RFC: <code>XAXX010101000</code>" in text
    assert "Código postal: <code>64000</code>" in text
    assert "Régimen" not in text


def test_split_message():
    text = "\n".join(f"línea {i:03d}" for i in range(100))
    chunks = split_message(text, limit=100)
    assert all(len(c) <= 100 for c in chunks)
    assert "\n".join(chunks) == text
    assert split_message("hola") == ["hola"]


def test_build_xlsx(db):
    e = add_expense(
        db,
        description="=HYPERLINK(\"http://malo\")",
        invoice_fields=[{"label": "Web ID", "value": "1"}],
        subtotal=211.64,
        iva=33.86,
    )
    wb = load_workbook(io.BytesIO(build_xlsx([e], TZ)))
    ws = wb["Gastos"]
    header = [c.value for c in ws[1]]
    row = {name: cell for name, cell in zip(header, ws[2])}
    assert ws.max_row == 2
    assert row["ID"].value == e.id
    assert row["Fecha"].value.date() == date(2026, 10, 3)
    assert row["Total"].value == 245.5
    assert row["Categoría"].value == "Supermercado"
    assert row["Factura"].value == "Pendiente"
    assert row["Datos para facturar"].value == "Web ID: 1"
    assert row["Registrado el"].value.hour == 12
    # El texto que parece fórmula se guarda como texto.
    assert row["Descripción"].data_type == "s"
    assert row["Descripción"].value == "=HYPERLINK(\"http://malo\")"
